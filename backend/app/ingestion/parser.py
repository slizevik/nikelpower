"""Extract text from PDF, DOCX, and PPTX (legacy) + facade полного пайплайна."""

from __future__ import annotations

from pathlib import Path

from app.ingestion.exceptions import SUPPORTED_DOCUMENT_EXTENSIONS, UnsupportedFormatError
from app.ingestion.language import detect_language_hint
from app.ingestion.parser_types import ParsedDocument

__all__ = [
    "ParsedDocument",
    "SUPPORTED_EXTENSIONS",
    "detect_language_hint",
    "parse_document",
    "parse_document_legacy",
    "iter_documents",
    "UnsupportedFormatError",
]

SUPPORTED_EXTENSIONS = SUPPORTED_DOCUMENT_EXTENSIONS


def parse_pdf(path: Path) -> ParsedDocument:
    import fitz

    pages: list[tuple[int, str]] = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                pages.append((i, text))
    full_text = "\n\n".join(t for _, t in pages)
    return ParsedDocument(
        source_path=str(path),
        doc_type="pdf",
        language_hint=detect_language_hint(full_text),
        pages=pages,
        full_text=full_text,
    )


def parse_docx(path: Path) -> ParsedDocument:
    from docx import Document

    doc = Document(path)
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    pages = [(i + 1, p) for i, p in enumerate(paragraphs)]
    full_text = "\n\n".join(paragraphs)
    return ParsedDocument(
        source_path=str(path),
        doc_type="docx",
        language_hint=detect_language_hint(full_text),
        pages=pages,
        full_text=full_text,
    )


def parse_pptx(path: Path) -> ParsedDocument:
    from pptx import Presentation

    prs = Presentation(path)
    pages: list[tuple[int, str]] = []
    for i, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"[notes] {notes}")
        if parts:
            pages.append((i, "\n".join(parts)))
    full_text = "\n\n".join(t for _, t in pages)
    return ParsedDocument(
        source_path=str(path),
        doc_type="pptx",
        language_hint=detect_language_hint(full_text),
        pages=pages,
        full_text=full_text,
    )


def parse_document(path: Path) -> ParsedDocument:
    """Полный пайплайн: LibreOffice → PDF → VLM для изображений."""
    from app.ingestion.orchestrator import parse_document_full, to_legacy_parsed_document

    return to_legacy_parsed_document(parse_document_full(path))


def parse_document_legacy(path: Path) -> ParsedDocument:
    """Быстрый парсинг без LibreOffice/VLM (только текст)."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    if suffix == ".pptx":
        return parse_pptx(path)
    raise UnsupportedFormatError(str(path), suffix)


def iter_documents(root: Path) -> list[Path]:
    if not root.exists():
        return []
    return [
        path
        for path in sorted(root.rglob("*"))
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    ]
